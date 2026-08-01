"""GX-3a: opt-in exact multi-row header carrier (XLSX).

GX-3a of ``FTR-XREF-AXIS-MAPPINGS-P4A2``. Proves the configured-depth,
merge-independent, lossless per-cell ``header_grid`` write/read path, and that
legacy behaviour is unchanged when the exact gate is not used. No grouped-XRef,
axis-mapping, or ``MultiIndex`` semantics are involved.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from spreadsheet_handling.io_backends.xlsx.openpyxl_parser import parse_workbook
from spreadsheet_handling.io_backends.xlsx.openpyxl_renderer import render_workbook
from spreadsheet_handling.io_backends.xlsx.xlsx_backend import ExcelBackend
from spreadsheet_handling.rendering.flow import build_render_plan
from spreadsheet_handling.rendering.ir import SheetIR, TableBlock, WorkbookIR

pytestmark = pytest.mark.ftr("FTR-XREF-AXIS-MAPPINGS-P4A2")


def _ir_with_grid(
    grid: tuple[tuple[str, ...], ...],
    data: list[list[str]],
    *,
    sheet: str = "mtx",
) -> WorkbookIR:
    depth = len(grid)
    n_cols = len(grid[0])
    tbl = TableBlock(
        frame_name=sheet,
        top=1,
        left=1,
        header_rows=depth,
        header_cols=1,
        n_rows=depth + len(data),
        n_cols=n_cols,
        headers=list(grid[-1]),
        header_map={h: i + 1 for i, h in enumerate(grid[-1])},
        data=data,
        header_grid=grid,
    )
    return WorkbookIR(sheets={sheet: SheetIR(name=sheet, tables=[tbl])})


def _write_ir(ir: WorkbookIR, path: Path) -> None:
    render_workbook(build_render_plan(ir), path)


def _roundtrip(
    grid: tuple[tuple[str, ...], ...],
    data: list[list[str]],
    tmp_path: Path,
    *,
    sheet: str = "mtx",
) -> TableBlock:
    path = tmp_path / "exact.xlsx"
    _write_ir(_ir_with_grid(grid, data, sheet=sheet), path)
    ir = parse_workbook(str(path), exact_header_depths={sheet: len(grid)})
    return ir.sheets[sheet].tables[0]


# --------------------------------------------------------------------------- #
# Write/read roundtrip                                                         #
# --------------------------------------------------------------------------- #

def test_arity2_no_merge_roundtrip_is_exact(tmp_path: Path) -> None:
    grid = (("", "Kredit", "Kredit", "Einlage"), ("row_id", "Ann", "Fix", "Gut"))
    data = [["r1", "a", "b", "c"], ["r2", "d", "e", "f"]]
    tbl = _roundtrip(grid, data, tmp_path)
    assert tbl.header_grid == grid
    assert tbl.header_rows == 2
    assert tbl.data == data


def test_arity3_roundtrip_is_exact(tmp_path: Path) -> None:
    grid = (
        ("", "Top", "Top", "Top"),
        ("", "Mid", "Mid", "Other"),
        ("row_id", "L1", "L2", "L3"),
    )
    data = [["r1", "1", "2", "3"]]
    tbl = _roundtrip(grid, data, tmp_path)
    assert tbl.header_grid == grid
    assert tbl.header_rows == 3
    assert tbl.data == data


def test_literal_delimiter_component_survives(tmp_path: Path) -> None:
    grid = (("Gr / oup", "Gr / oup"), ("a.b", "leaf"))
    tbl = _roundtrip(grid, [["x", "y"]], tmp_path)
    assert tbl.header_grid == grid
    assert tbl.header_grid[0][0] == "Gr / oup"


def test_blank_physical_cells_retained_without_filtering(tmp_path: Path) -> None:
    grid = (("", "A", ""), ("row_id", "leaf", "leaf2"))
    tbl = _roundtrip(grid, [["r1", "1", "2"]], tmp_path)
    # Blanks kept verbatim (no forward-fill, no drop).
    assert tbl.header_grid == grid


def test_deterministic_data_start_below_configured_depth(tmp_path: Path) -> None:
    grid = (("", "A", "B"), ("row_id", "x", "y"))
    data = [["r1", "1", "2"], ["r2", "3", "4"]]
    path = tmp_path / "exact.xlsx"
    _write_ir(_ir_with_grid(grid, data), path)
    wb = openpyxl.load_workbook(str(path))
    ws = wb["mtx"]
    # Header occupies rows 1-2; data starts at row 3.
    assert ws.cell(row=3, column=1).value == "r1"
    assert ws.cell(row=2, column=1).value == "row_id"
    tbl = parse_workbook(str(path), exact_header_depths={"mtx": 2}).sheets["mtx"].tables[0]
    assert tbl.data == data


# --------------------------------------------------------------------------- #
# Merge independence                                                          #
# --------------------------------------------------------------------------- #

def _author_2row(path: Path, *, merged: bool) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "mtx"
    if merged:
        ws["B1"] = "Kredit"
        ws.merge_cells("B1:C1")
    else:
        ws["B1"] = "Kredit"
        ws["C1"] = "Kredit"
    ws["D1"] = "Einlage"
    ws["A2"], ws["B2"], ws["C2"], ws["D2"] = "row_id", "Ann", "Fix", "Gut"
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "r1", "a", "b", "c"
    wb.save(str(path))


def test_configured_depth_with_merges_resolves_masters(tmp_path: Path) -> None:
    path = tmp_path / "merged.xlsx"
    _author_2row(path, merged=True)
    tbl = parse_workbook(str(path), exact_header_depths={"mtx": 2}).sheets["mtx"].tables[0]
    assert tbl.header_grid == (
        ("", "Kredit", "Kredit", "Einlage"),
        ("row_id", "Ann", "Fix", "Gut"),
    )
    assert tbl.n_cols == 4  # leaf-row extent, unaffected by the blank A1


def test_user_unmerged_equivalent_reconstructs_identically(tmp_path: Path) -> None:
    merged = tmp_path / "merged.xlsx"
    unmerged = tmp_path / "unmerged.xlsx"
    _author_2row(merged, merged=True)
    _author_2row(unmerged, merged=False)
    g_merged = parse_workbook(str(merged), exact_header_depths={"mtx": 2}).sheets["mtx"].tables[0].header_grid
    g_unmerged = parse_workbook(str(unmerged), exact_header_depths={"mtx": 2}).sheets["mtx"].tables[0].header_grid
    assert g_merged == g_unmerged


# --------------------------------------------------------------------------- #
# Gate / malformed input                                                       #
# --------------------------------------------------------------------------- #

@pytest.mark.parametrize("bad", [0, -1, "2", 2.0, None])
def test_malformed_configured_depth_is_rejected(tmp_path: Path, bad: object) -> None:
    path = tmp_path / "exact.xlsx"
    _write_ir(_ir_with_grid((("", "A"), ("row_id", "x")), [["r1", "1"]]), path)
    with pytest.raises(ValueError, match="positive integer"):
        parse_workbook(str(path), exact_header_depths={"mtx": bad})  # type: ignore[dict-item]


def test_malformed_depth_fails_before_producing_ir(tmp_path: Path) -> None:
    # Atomicity: validation happens before any workbook is opened/parsed.
    with pytest.raises(ValueError, match="positive integer"):
        parse_workbook(str(tmp_path / "does_not_exist.xlsx"), exact_header_depths={"s": 0})


# --------------------------------------------------------------------------- #
# Legacy no-regression (same primitives, exact gate unused)                    #
# --------------------------------------------------------------------------- #

def test_legacy_read_of_exact_file_has_no_header_grid(tmp_path: Path) -> None:
    grid = (("", "Kredit", "Kredit", "Einlage"), ("row_id", "Ann", "Fix", "Gut"))
    path = tmp_path / "exact.xlsx"
    _write_ir(_ir_with_grid(grid, [["r1", "a", "b", "c"]]), path)
    # Without the exact gate, the table carries no header_grid (legacy path).
    tbl = parse_workbook(str(path)).sheets["mtx"].tables[0]
    assert tbl.header_grid is None


def test_ordinary_single_row_sheet_unchanged_without_gate(tmp_path: Path) -> None:
    import pandas as pd

    frames = {"data": pd.DataFrame({"id": [1, 2], "name": ["a", "b"]})}
    path = tmp_path / "plain.xlsx"
    ExcelBackend().write_multi(frames, str(path))
    back = ExcelBackend().read_multi(str(path), header_levels=1)
    assert list(back["data"].columns) == ["id", "name"]
    assert not hasattr(back["data"].columns, "levels") or not isinstance(
        back["data"].columns, pd.MultiIndex
    )


def test_backend_option_gate_reaches_parser(tmp_path: Path) -> None:
    grid = (("", "A", "B"), ("row_id", "x", "y"))
    path = tmp_path / "exact.xlsx"
    _write_ir(_ir_with_grid(grid, [["r1", "1", "2"]]), path)
    # The read option threads through to the parser (frame projection itself is
    # legacy/GX-3b; here we only assert the gate is reachable and does not error).
    frames = ExcelBackend().read_multi(
        str(path),
        header_levels=2,
        options={"extra": {"exact_header_depths": {"mtx": 2}}},
    )
    assert "mtx" in frames
